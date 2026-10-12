"""Meeting record: autosave while recording, and the final transcript / summary files.

No audio is ever written — only text.
"""
import json
import os
import re
import time

from mk.transcriber import MODELS_DIR

APP_DIR = os.path.dirname(MODELS_DIR)
SETTINGS_FILE = os.path.join(APP_DIR, "data", "settings.json")
DEFAULT_OUTPUT_DIR = os.path.join(APP_DIR, "meetings")


def load_settings():
    try:
        with open(SETTINGS_FILE, encoding="utf-8") as f:
            return json.load(f)
    except (OSError, ValueError):
        return {}


def save_settings(settings):
    os.makedirs(os.path.dirname(SETTINGS_FILE), exist_ok=True)
    with open(SETTINGS_FILE, "w", encoding="utf-8") as f:
        json.dump(settings, f, ensure_ascii=False, indent=1)


def _safe(name):
    return re.sub(r'[\\/:*?"<>|]+', " ", name).strip()[:60]


def _next_prefix(output_dir, date):
    """Next free "<date>_NN" in output_dir (folders may carry a title after the number)."""
    used = set()
    if os.path.isdir(output_dir):
        for name in os.listdir(output_dir):
            m = re.match(rf"^{date}_(\d{{2,}})", name)
            if m:
                used.add(int(m.group(1)))
    n = 1
    while n in used:
        n += 1
    return f"{date}_{n:02d}"


def _write_text(path, text, hidden=False):
    """Atomic write. Notepad-friendly: UTF-8 with BOM and CRLF line endings."""
    tmp = path + ".tmp"
    with open(tmp, "w", encoding="utf-8-sig", newline="\r\n") as f:
        f.write(text)
    os.replace(tmp, path)
    if hidden and os.name == "nt":
        import ctypes
        ctypes.windll.kernel32.SetFileAttributesW(path, 0x2)   # FILE_ATTRIBUTE_HIDDEN


def _plain(markdown):
    """Summary markdown → plain text that reads well in Notepad or pasted into chat/email."""
    out = []
    for line in markdown.strip().splitlines():
        line = line.rstrip().replace("**", "")
        if line.startswith("#"):
            if out and out[-1]:
                out.append("")
            out.append("■ " + line.lstrip("#").strip())
        else:
            out.append(re.sub(r"^(\s*)[-*]\s+", r"\1• ", line))
    return "\n".join(out)


def _hms(ts):
    return time.strftime("%H:%M:%S", time.localtime(ts))


class Meeting:
    """lines: [{"time": epoch, "speaker": label, "me": bool, "en": str, "ko": str | None}]"""

    def __init__(self, output_dir):
        self.started = time.time()
        self.ended = None
        self.title = ""
        self.lines = []
        self.summary = None
        # "20261012_01": date + the meeting's number that day. Used for the folder and every file name.
        self.prefix = _next_prefix(output_dir, time.strftime("%Y%m%d", time.localtime(self.started)))
        self.folder = os.path.join(output_dir, self.prefix)

    # --- metadata ------------------------------------------------------------
    def participants(self):
        seen = []
        for ln in self.lines:
            if ln["speaker"] not in seen:
                seen.append(ln["speaker"])
        return seen

    def period(self):
        end = self.ended or (self.lines[-1]["time"] if self.lines else self.started)
        minutes = max(1, round((end - self.started) / 60))
        return (f"{time.strftime('%Y-%m-%d %H:%M', time.localtime(self.started))} ~ "
                f"{time.strftime('%H:%M', time.localtime(end))} ({minutes}분)")

    def english_transcript(self):
        """Plain transcript used as summary input: consecutive lines of one speaker are joined."""
        out = []
        for ln in self.lines:
            if out and out[-1][0] == ln["speaker"]:
                out[-1][1].append(ln["en"])
            else:
                out.append((ln["speaker"], [ln["en"]]))
        return "\n".join(f"{spk}: {' '.join(texts)}" for spk, texts in out)

    def set_title(self, title):
        """Sets the title and renames the folder to include it."""
        self.title = title.strip()
        name = self.prefix + (f"_{_safe(self.title)}" if self.title else "")
        new = os.path.join(os.path.dirname(self.folder), name)
        if new != self.folder:
            if os.path.exists(self.folder) and not os.path.exists(new):
                os.rename(self.folder, new)
            self.folder = new

    # --- files ---------------------------------------------------------------
    def path(self, kind, ext):
        return os.path.join(self.folder, f"{self.prefix}_{kind}.{ext}")

    def autosave(self):
        """Hidden working files, rewritten while the meeting runs: the data (for re-opening later) and a
        plain-text transcript that survives a crash before the meeting is saved."""
        os.makedirs(self.folder, exist_ok=True)
        d = {"title": self.title, "started": self.started, "ended": self.ended,
             "lines": self.lines, "summary": self.summary}
        _write_text(self.path("data", "json"), json.dumps(d, ensure_ascii=False, indent=1), hidden=True)
        _write_text(self.path("전문_자동저장", "txt"), self._transcript_text(), hidden=True)

    def _header(self, kind):
        return (f"{kind} — {self.title or '회의'}\n"
                f"일시: {self.period()}\n"
                f"참석자: {', '.join(self.participants())}\n")

    def _transcript_text(self):
        parts = [self._header("회의 전문")]
        for ln in self.lines:
            parts.append(f"[{_hms(ln['time'])}] {ln['speaker']}\n{ln['en']}\n{ln['ko'] or '(번역 없음)'}\n")
        return "\n".join(parts)

    def save_final(self):
        """Writes 회의록.docx (summary + transcript) and 요약.txt (if a summary exists). Returns the folder."""
        self.autosave()
        if self.summary:
            _write_text(self.path("요약", "txt"), self._header("회의 요약") + "\n" + _plain(self.summary) + "\n")
        self._write_docx()
        return self.folder

    def _write_docx(self):
        from docx import Document
        from docx.shared import Pt, RGBColor

        doc = Document()
        style = doc.styles["Normal"]
        style.font.name = "맑은 고딕"
        style.font.size = Pt(10.5)
        style.element.rPr.rFonts.set(
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia", "맑은 고딕")

        doc.add_heading(self.title or "회의록", level=0)
        doc.add_paragraph(f"일시: {self.period()}\n참석자: {', '.join(self.participants())}")

        if self.summary:
            doc.add_heading("요약", level=1)
            for raw in self.summary.strip().splitlines():
                line = raw.rstrip()
                if not line:
                    continue
                if line.startswith("#"):
                    doc.add_heading(line.lstrip("#").strip(), level=2)
                elif re.match(r"^\s*[-*•]\s+", line):
                    doc.add_paragraph(re.sub(r"^\s*[-*•]\s+", "", line).replace("**", ""), style="List Bullet")
                else:
                    doc.add_paragraph(line.replace("**", ""))
            doc.add_page_break()

        doc.add_heading("전문", level=1)
        for ln in self.lines:
            p = doc.add_paragraph()
            r = p.add_run(f"[{_hms(ln['time'])}] {ln['speaker']}")
            r.bold = True
            p.add_run("\n" + ln["en"])
            ko = p.add_run("\n" + (ln["ko"] or "(번역 없음)"))
            ko.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        doc.save(self.path("회의록", "docx"))
